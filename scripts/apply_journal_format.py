# -*- coding: utf-8 -*-
"""《计算机应用与软件》格式后处理：① 表格改网格型(外框 1 磅、内线 1/2 磅)；② 页眉(奇/偶不同、首页两行)。
   用法：python apply_journal_format.py   (在 build_paper.py 之后运行；会就地更新投稿稿 DOCX)
"""
import copy
import re as _re
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
R = Path(str(Path(__file__).resolve().parents[1]))
DOC = R / "基于配对跨域协议的微调收益与代价测量_投稿稿.docx"
TITLE = "基于配对跨域协议的微调收益与代价测量"
FIRST_AUTHOR = "常开开"          # 待作者信息确定后替换
JOURNAL = "计算机应用与软件"


def set_borders(tbl, outer_sz=8, inner_sz=4):
    """给表格加网格边框：外框 1 磅(sz=8)、内部 0.5 磅(sz=4)。"""
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for tag, sz in (("top", outer_sz), ("left", outer_sz), ("bottom", outer_sz), ("right", outer_sz),
                    ("insideH", inner_sz), ("insideV", inner_sz)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def enable_even_odd(doc):
    st = doc.settings.element
    if st.find(qn("w:evenAndOddHeaders")) is None:
        e = OxmlElement("w:evenAndOddHeaders")
        st.append(e)


def put_header(hdr, lines, size=8):
    """写入页眉段落(逐行居中，指定字号)。"""
    if hdr is None:
        return
    hdr.is_linked_to_previous = False
    for p in list(hdr.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    first = hdr.paragraphs[0]
    first.text = lines[0] if lines else ""
    for extra in lines[1:]:
        first.add_run().add_break()
        first.add_run(extra)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    for p in hdr.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(size)


def number_formulas(doc):
    """给独立公式段加全文统一编号 (1)(2)…(右对齐制表位)。"""
    from docx.shared import Pt as _Pt
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement as _El
    n = 0
    for p in doc.paragraphs:
        el = p._element
        has_math = el.findall(".//" + _qn("m:oMath")) or el.findall(".//" + _qn("m:oMathPara"))
        txt = (p.text or "").strip()
        if not has_math and "$$" not in txt:
            continue
        if len(txt) > 400:
            continue
        # 清除旧编号（避免重复 (1)(1)）
        for _r in list(p.runs):
            if re.fullmatch(r"\s*\(\d+\)", _r.text or ""):
                _r._element.getparent().remove(_r._element)
        n += 1
        pPr = el.get_or_add_pPr()
        # 清除旧制表位，避免与 4800 冲突
        _old_tabs = pPr.find(_qn("w:tabs"))
        if _old_tabs is not None:
            pPr.remove(_old_tabs)
        tabs = _El("w:tabs")
        tab = _El("w:tab")
        tab.set(_qn("w:val"), "right")
        tab.set(_qn("w:pos"), "4800")
        tabs.append(tab)
        pPr.append(tabs)
        run = p.add_run("\t(%d)" % n)
        run.font.size = _Pt(10.5)
    return n



TEMPLATE_MARGINS = (3.20, 1.30, 1.90, 1.30)   # 上/下/左/右（cm，取自官方模板）


def apply_template(doc):
    """① 页面边距按官方模板；② 标题用 Heading 1/2，正文用模板的正文样式。"""
    import re as _re
    from docx.shared import Cm
    from docx.enum.style import WD_STYLE_TYPE
    for sec in doc.sections:
        sec.top_margin, sec.bottom_margin = Cm(TEMPLATE_MARGINS[0]), Cm(TEMPLATE_MARGINS[1])
        sec.left_margin, sec.right_margin = Cm(TEMPLATE_MARGINS[2]), Cm(TEMPLATE_MARGINS[3])
    names = {st.name for st in doc.styles if st.type == WD_STYLE_TYPE.PARAGRAPH}
    body_style = next((n for n in ("正文文本首行缩进", "正文缩进", "Body Text") if n in names), None)
    h1, h2 = ("Heading 1" if "Heading 1" in names else None), ("Heading 2" if "Heading 2" in names else None)
    n_h1 = n_h2 = n_body = 0
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t or len(t) > 60:
            continue
        if _re.match(r"^(0 引|1 相关|2 配对|3 结果|4 讨论|5 结论|摘要与关键词|参考文献|文末另页)", t):
            if h1: p.style = doc.styles[h1]; n_h1 += 1
        elif _re.match(r"^\d+\.\d+\s", t):
            if h2: p.style = doc.styles[h2]; n_h2 += 1
        elif body_style and not _re.match(r"^(图|表|注|\[)", t):
            p.style = doc.styles[body_style]; n_body += 1
    return n_h1, n_h2, n_body, body_style


def fix_columns(doc):
    """显式重建三节：前置(1栏,左1.90) → 引言起正文(2栏,左1.90) → 文末(1栏,左1.75)。"""
    from docx.oxml.ns import qn as Q
    from lxml import etree as ET
    CM = 567

    def clear_breaks():
        n = 0
        for ch in doc.element.body.iterchildren():
            if ch.tag == Q("w:p"):
                sp = ch.find(Q("w:pPr") + "/" + Q("w:sectPr"))
                if sp is not None:
                    ch.find(Q("w:pPr")).remove(sp); n += 1
        return n

    def set_pg(sect, left_cm):
        pg = sect.find(Q("w:pgMar"))
        if pg is None:
            pg = ET.SubElement(sect, Q("w:pgMar"))
        pg.set(Q("w:top"), str(int(3.20 * CM))); pg.set(Q("w:bottom"), str(int(1.30 * CM)))
        pg.set(Q("w:left"), str(int(left_cm * CM))); pg.set(Q("w:right"), str(int(1.30 * CM)))

    def break_at(idx, cols, left_cm):
        el = ET.fromstring(ET.tostring(base))
        c = el.find(Q("w:cols"))
        if c is None:
            c = ET.SubElement(el, Q("w:cols"))
        c.set(Q("w:num"), str(cols))
        set_pg(el, left_cm)
        p = doc.paragraphs[idx - 1]._element
        pPr = p.find(Q("w:pPr"))
        if pPr is None:
            pPr = ET.SubElement(p, Q("w:pPr")); p.insert(0, pPr)
        pPr.append(el)

    ps = doc.paragraphs
    iy = next((k for k, p in enumerate(ps) if re.match(r"^0\s*引", (p.text or "").strip())), None)
    iw = next((k for k, p in enumerate(ps) if (p.text or "").strip().startswith("文末另页")), None)
    base = ET.fromstring(ET.tostring(doc.sections[-1]._sectPr))   # 干净基准
    cleared = clear_breaks()
    if iy:
        break_at(iy, 1, 1.90)      # 前置：单栏
    if iw:
        break_at(iw, 2, 1.90)      # 正文：双栏
    # 末节（文末）最后设置，避免被继承污染
    last = doc.sections[-1]._sectPr
    c = last.find(Q("w:cols"))
    if c is None:
        c = ET.SubElement(last, Q("w:cols"))
    c.set(Q("w:num"), "1")
    set_pg(last, 1.75)
    return cleared + (1 if iy else 0) + (1 if iw else 0)


def main():
    if not DOC.exists():
        print("  ✗ 未找到 DOCX"); return
    doc = Document(str(DOC))

    # ① 网格表
    for i, tbl in enumerate(doc.tables, 1):
        set_borders(tbl)
    print("  ① 已把 %d 张表改为网格型(外框 1 磅 / 内线 0.5 磅)✓" % len(doc.tables))

    # ② 页眉
    enable_even_odd(doc)
    odd_txt = "%s，等：%s" % (FIRST_AUTHOR, TITLE)
    for sec in doc.sections:
        sec.different_first_page_header_footer = True                 # 首页独立页眉
        put_header(sec.first_page_header, ["计算机应用与软件", "Computer Applications and Software"])
        put_header(sec.header, [odd_txt])                             # 奇数页
        try:
            put_header(sec.even_page_header, ["计算机应用与软件"])       # 偶数页
        except Exception as exc:                                       # noqa: BLE001
            print("    偶页页眉跳过：", str(exc)[:60])
    print("  ② 页眉：首页两行「计算机应用与软件 / Computer Applications and Software」；"
          "奇数页「%s」；偶数页「计算机应用与软件」✓" % odd_txt)

    # ③ 页脚清空
    for s in doc.sections:
        for hdr in (s.footer, s.even_page_footer, s.first_page_footer):
            try:
                hdr.is_linked_to_previous = False
                for p in hdr.paragraphs:
                    p.text = ""
            except Exception:
                pass
    print("  ③ 页脚已清空 ✓")

    a1, a2, a3, bs = apply_template(doc)
    print("  ⑤ 模板化：边距 3.20/1.30/1.90/1.30 cm；Heading1 %d 段、Heading2 %d 段、正文样式 %s %d 段 ✓" % (a1, a2, bs, a3))
    ncol = fix_columns(doc)
    print("  ⑥ 分栏：前置单栏 / 引言起双栏 / 文末单栏(左 1.75cm) ✓（调整 %d 次）" % ncol)
    ncol = fix_columns(doc)
    print("  ⑥ 分栏：前置单栏、自“0 引　言”起双栏 ✓（调整 %d 次）" % ncol)
    # 删除改稿遗留的空壳段（只含 ** / * 的段落）
    _junk = 0
    for _p in list(doc.paragraphs):
        if (_p.text or "").strip() in ("**", "*", "***"):
            _p._element.getparent().remove(_p._element)
            _junk += 1
    print("  ⑧ 已删除空壳段 %d 个 ✓" % _junk)
    nf = number_formulas(doc)
    print("  ④ 已为 %d 个独立公式加编号 (1)(2)… ✓" % nf)
    doc.save(str(DOC))
    print("  已保存：", DOC.name)


if __name__ == "__main__":
    main()
